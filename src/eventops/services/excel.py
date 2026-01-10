from __future__ import annotations

from collections.abc import Sequence
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook


@dataclass(frozen=True)
class ExcelSummary:
    file: str
    sheet_names: Sequence[str]
    sheet_used: str
    rows: int
    cols: int
    columns: Sequence[str]
    preview: Sequence[Sequence[str]]  # rows of stringified values


class ExcelError(Exception):
    """Errores controlados al procesar Excel."""


def summarize_excel(
    file_path: Path,
    sheet: str | None = None,
    preview_rows: int = 5,
) -> ExcelSummary:
    """
    Lee un .xlsx y devuelve un resumen usable por la CLI.
    - No imprime nada.
    - Lanza ExcelError con mensajes claros para el usuario.
    """
    if not file_path.exists():
        raise ExcelError(f"El archivo no existe: {file_path}")

    if file_path.suffix.lower() != ".xlsx":
        raise ExcelError("Formato no soportado. Solo se admite .xlsx en el Día 3.")

    try:
        wb = load_workbook(filename=str(file_path), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ExcelError(f"No se pudo abrir el Excel: {exc}") from exc

    sheet_names = wb.sheetnames
    if not sheet_names:
        raise ExcelError("El archivo Excel no contiene hojas.")

    if sheet is None:
        sheet_used = sheet_names[0]
    else:
        if sheet not in sheet_names:
            raise ExcelError(
                f"La hoja '{sheet}' no existe. Hojas disponibles: {', '.join(sheet_names)}"
            )
        sheet_used = sheet

    ws = wb[sheet_used]

    # Detectamos rango real usando la fila 1 como cabecera.
    # openpyxl devuelve None en celdas vacías. Convertimos a str limpio.
    header = []
    for cell in ws[1]:
        header.append("" if cell.value is None else str(cell.value).strip())

    # Si la cabecera está vacía, seguimos igualmente, pero lo reflejamos.
    columns = header

    # Contamos filas/cols según max_row/max_column (aprox, pero útil para resumen).
    rows = ws.max_row or 0
    cols = ws.max_column or 0

    # Preview: filas 1..preview_rows+1 (incluye cabecera + primeras filas)
    preview_data: list[list[str]] = []
    max_preview = min(rows, preview_rows + 1)

    for r in range(1, max_preview + 1):
        row_vals: list[str] = []
        for c in range(1, cols + 1):
            v = ws.cell(row=r, column=c).value
            row_vals.append("" if v is None else str(v))
        preview_data.append(row_vals)

    return ExcelSummary(
        file=str(file_path),
        sheet_names=sheet_names,
        sheet_used=sheet_used,
        rows=rows,
        cols=cols,
        columns=columns,
        preview=preview_data,
    )
